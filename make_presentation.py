"""
Генерация презентации по мини-проекту «Оценка кредитного риска».
Создаёт credit_scoring_presentation.pptx с 7 слайдами.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ---------- Данные ----------
AUTHOR = "Козаченко Сергей Александрович"
GROUP = "МИБСИ261"
TEACHER = "Лебедев Олег Борисович"
YEAR = "2026"
GITHUB = "https://github.com/molot88888888-spec/credit-scoring"

# ---------- Пути ----------
BASE = os.path.dirname(os.path.abspath(__file__))
NOTEBOOKS = os.path.join(BASE, "notebooks")


def img(name, folder=NOTEBOOKS):
    """Возвращает путь к файлу, если он существует, иначе None."""
    path = os.path.join(folder, name)
    return path if os.path.exists(path) else None


# ---------- Создаём презентацию ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]  # пустой макет


def add_title(slide, text, size=32):
    """Заголовок слайда."""
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3),
                                  Inches(12.3), Inches(0.9))
    tf = tb.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def add_text(slide, text, left=0.6, top=1.4, width=12.1, height=5.5, size=18):
    """Текстовый блок с переносами."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_image(slide, path, left, top, width):
    """Вставляет картинку, если файл существует."""
    if path:
        slide.shapes.add_picture(path, Inches(left), Inches(top),
                                 width=Inches(width))


# ============================================================
# СЛАЙД 1. Титул
# ============================================================
s = prs.slides.add_slide(BLANK)

tb = s.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11.3), Inches(1.5))
tf = tb.text_frame
tf.text = "Оценка кредитного риска"
p = tf.paragraphs[0]
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
p.alignment = PP_ALIGN.CENTER

tb2 = s.shapes.add_textbox(Inches(1), Inches(3.4), Inches(11.3), Inches(0.8))
tf2 = tb2.text_frame
tf2.text = "Построение скоринговой модели"
p2 = tf2.paragraphs[0]
p2.font.size = Pt(24)
p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
p2.alignment = PP_ALIGN.CENTER

tb3 = s.shapes.add_textbox(Inches(1), Inches(4.8), Inches(11.3), Inches(2))
tf3 = tb3.text_frame
tf3.text = (
    f"Выполнил: {AUTHOR}\n"
    f"Группа: {GROUP}\n"
    f"Проверил: {TEACHER}\n"
    f"Год: {YEAR}"
)
for i, line in enumerate(tf3.text.split("\n")):
    p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
    p.text = line
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.CENTER

tb4 = s.shapes.add_textbox(Inches(1), Inches(6.6), Inches(11.3), Inches(0.5))
tf4 = tb4.text_frame
tf4.text = GITHUB
p4 = tf4.paragraphs[0]
p4.font.size = Pt(12)
p4.font.color.rgb = RGBColor(0x00, 0x7B, 0xFF)
p4.alignment = PP_ALIGN.CENTER

# ============================================================
# СЛАЙД 2. Постановка задачи
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Постановка задачи")
add_text(s, (
    "Что делаем:\n"
    "•  Бинарная классификация: предсказать вероятность дефолта заёмщика\n"
    "•  Вход: 23 признака (лимит, возраст, история платежей, суммы счетов)\n"
    "•  Выход: вероятность дефолта → скоринговый балл 300–850\n"
    "\n"
    "Почему это важно:\n"
    "•  Удержать клиента дешевле, чем привлечь нового\n"
    "•  Банк должен оценить риск до выдачи кредита\n"
    "\n"
    "Ключевые особенности:\n"
    "•  Дисбаланс классов: дефолтов ~22%\n"
    "•  Интерпретируемость: решения должны быть объяснимы\n"
    "•  Асимметрия ошибок: выдать кредит мошеннику дороже, чем отказать надёжному\n"
    "\n"
    "Метрики: ROC-AUC, recall, F1, precision (не accuracy!)"
), size=16)

# ============================================================
# СЛАЙД 3. Данные и EDA
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Данные и разведочный анализ")

add_text(s, (
    "Датасет:\n"
    "•  Default of Credit Card Clients (UCI, Тайвань, 2005)\n"
    "•  30 000 клиентов, 23 признака\n"
    "•  Метка: 0 — вернул, 1 — дефолт\n"
    "\n"
    "Что показал анализ:\n"
    "•  Дисбаланс классов: 22% дефолтов\n"
    "•  PAY_0 — главный признак:\n"
    "     PAY_0 = -2 (нет долга) → 14% дефолтов\n"
    "     PAY_0 = 2 (просрочка 2 мес.) → 69% дефолтов\n"
    "•  Слабые признаки: EDUCATION, SEX, MARRIAGE"
), left=0.6, top=1.4, width=6.0, height=5.5, size=15)

add_image(s, img("class_balance.png"), left=6.8, top=1.6, width=6.0)
add_image(s, img("pay0_default_rate.png"), left=6.8, top=4.2, width=6.0)

# ============================================================
# СЛАЙД 4. Модели и метрики
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Модели и метрики")

add_text(s, (
    "Обучили 4 модели:\n"
    "\n"
    "LogisticRegression     —  Acc 0.684  Prec 0.373  Rec 0.628  F1 0.468  AUC 0.716\n"
    "DecisionTree            —  Acc 0.781  Prec 0.505  Rec 0.540  F1 0.522  AUC 0.756\n"
    "RandomForest            —  Acc 0.802  Prec 0.561  Rec 0.483  F1 0.519  AUC 0.762\n"
    "GradientBoosting        —  Acc 0.816  Prec 0.657  Rec 0.357  F1 0.463  AUC 0.780\n"
    "\n"
    "Что важно:\n"
    "•  GradientBoosting — лучший по ROC-AUC (0.780)\n"
    "•  LogisticRegression — лучший по recall (0.628) и интерпретируем\n"
    "•  Для банка критичен recall: пропустить дефолт дороже,\n"
    "   чем отказать надёжному клиенту"
), left=0.6, top=1.4, width=7.0, height=5.5, size=13)

add_image(s, img("roc_curves.png", folder=BASE), left=8.0, top=1.8, width=5.0)

# ============================================================
# СЛАЙД 5. Интерпретация (SHAP)
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Интерпретация: какие признаки влияют на риск")

add_text(s, (
    "Топ-5 признаков (feature importance):\n"
    "1. PAY_0 — история платежей (сентябрь)\n"
    "2. PAY_2 — август\n"
    "3. PAY_3 — июль\n"
    "4. LIMIT_BAL — лимит по карте\n"
    "5. PAY_AMT1 — сумма последнего платежа\n"
    "\n"
    "SHAP-анализ подтвердил:\n"
    "•  Высокий PAY_0 (просрочка) → ↑↑ вероятность дефолта\n"
    "•  Высокий LIMIT_BAL → ↓ вероятность дефолта\n"
    "•  SEX, MARRIAGE, EDUCATION — почти не влияют\n"
    "\n"
    "Практический вывод:\n"
    "Модель можно упростить, исключив демографические признаки —\n"
    "это снизит риск дискриминации и не ухудшит качество."
), left=0.6, top=1.4, width=6.2, height=5.5, size=14)

add_image(s, img("feature_importance.png"), left=7.0, top=1.5, width=5.8)
add_image(s, img("shap_summary.png"), left=7.0, top=4.3, width=5.8)

# ============================================================
# СЛАЙД 6. Демо веб-сервиса
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Работающий веб-сервис на Flask")

add_text(s, (
    "Что сделано:\n"
    "•  Flask-приложение с HTML-формой\n"
    "•  Все 23 признака с подсказками\n"
    "•  Скоринговый балл 300–850\n"
    "•  Решение: Одобрить / Одобрить с осторожностью / Отказать\n"
    "\n"
    "Пример работы:\n"
    "\n"
    "Надёжный клиент:\n"
    "•  LIMIT_BAL = 500 000, AGE = 35, все PAY = -1\n"
    "•  Вероятность дефолта: 8%\n"
    "•  Скоринговый балл: 803\n"
    "•  Решение: Одобрить\n"
    "\n"
    "Рискованный клиент:\n"
    "•  LIMIT_BAL = 20 000, AGE = 22, PAY_0 = 2\n"
    "•  Вероятность дефолта: 74%\n"
    "•  Скоринговый балл: 442\n"
    "•  Решение: Отказать\n"
    "\n"
    f"Код: {GITHUB}"
), left=0.6, top=1.4, width=7.0, height=5.5, size=13)

# Если есть скриншот формы — вставим
add_image(s, img("web_form.png"), left=8.0, top=1.8, width=5.0)

# ============================================================
# СЛАЙД 7. Выводы и ограничения
# ============================================================
s = prs.slides.add_slide(BLANK)
add_title(s, "Выводы и ограничения")

add_text(s, (
    "Что сделано:\n"
    "•  Загружен и проанализирован датасет (30 000 клиентов)\n"
    "•  Обучены 4 модели классификации\n"
    "•  Проведена интерпретация (feature importance + SHAP)\n"
    "•  Создан веб-сервис со скоринговым баллом\n"
    "•  Код выложен на GitHub\n"
    "\n"
    "Выводы:\n"
    "1.  История платежей (PAY_*) — главный фактор риска дефолта\n"
    "2.  Демографические признаки почти не влияют — их можно исключить\n"
    "3.  LogisticRegression — оптимальна для скоринга:\n"
    "    интерпретируема и ловит больше дефолтов\n"
    "4.  GradientBoosting — лучший по ROC-AUC, но менее объясним\n"
    "\n"
    "Ограничения:\n"
    "•  Датасет из одной страны и одного года (Тайвань, 2005)\n"
    "•  Нет внешних данных (кредитная история, макроэкономика)\n"
    "•  Модель не учитывает временную динамику\n"
    "\n"
    "Что дальше:\n"
    "•  Добавить SHAP-объяснения в веб-интерфейс\n"
    "•  Применить SMOTE для улучшения recall\n"
    "•  Развернуть на VPS с Nginx + Gunicorn"
), size=13)


# ---------- Сохранение ----------
output = os.path.join(BASE, "credit_scoring_presentation.pptx")
prs.save(output)
print(f"Готово! Файл сохранён: {output}")
print(f"Слайдов: {len(prs.slides)}")